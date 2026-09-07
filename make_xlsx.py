# -*- coding: utf-8 -*-
"""Generate Login_Registration_Test_Suite.xlsx - Manual Testing portfolio workbook."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = r"D:\สหกิจ\Login-Registration-Testing\Login_Registration_Test_Suite.xlsx"

HDR_FILL = PatternFill("solid", fgColor="1F4E79")
HDR_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, size=14, color="1F4E79")
WRAP = Alignment(vertical="top", wrap_text=True)
CENTER = Alignment(vertical="top", horizontal="center", wrap_text=True)
THIN = Border(*[Side(style="thin", color="B0B0B0")] * 4)
PASS_FILL = PatternFill("solid", fgColor="C6EFCE")
FAIL_FILL = PatternFill("solid", fgColor="FFC7CE")
SEV_FILL = {"High": PatternFill("solid", fgColor="FFC7CE"),
            "Medium": PatternFill("solid", fgColor="FFEB9C"),
            "Low": PatternFill("solid", fgColor="DDEBF7")}

wb = Workbook()

def style_header(ws, ncols, row=1):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = CENTER
        cell.border = THIN

def finish(ws, ncols, nrows, widths, title=None, title_row=None):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r in range((title_row or 1), nrows + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = THIN
            if cell.alignment is None or not cell.alignment.wrapText:
                cell.alignment = WRAP
    ws.freeze_panes = ws.cell(row=(title_row or 1) + 1, column=1)

# ---------- Sheet 1: Test Scenarios ----------
ws = wb.active
ws.title = "Test Scenarios"
ws["A1"] = "Login & Registration Testing — Test Scenarios"
ws["A1"].font = TITLE_FONT
headers = ["Scenario ID", "Module", "Test Scenario", "Type", "Requirement"]
ws.append([])
ws.append(headers)
style_header(ws, len(headers), row=3)
scenarios = [
    ("TS-01", "Registration", "สมัครสมาชิกด้วยข้อมูลถูกต้องทุกช่อง", "Positive", "R-01"),
    ("TS-02", "Registration", "สมัครสมาชิกโดยไม่กรอก Email", "Negative", "R-02"),
    ("TS-03", "Registration", "สมัครสมาชิกด้วย Email ผิดรูปแบบ (ไม่มี @)", "Negative", "R-03"),
    ("TS-04", "Registration", "สมัครสมาชิกที่ Password ไม่ตรงกับ Confirm Password", "Negative", "R-05"),
    ("TS-05", "Registration", "สมัครสมาชิกด้วย Email ที่มีในระบบแล้ว", "Negative", "R-06"),
    ("TS-06", "Login", "Login ด้วย Email และ Password ที่ถูกต้อง", "Positive", "R-07"),
    ("TS-07", "Login", "Login ด้วย Password ที่ผิด", "Negative", "R-08"),
    ("TS-08", "Login", "Login โดยไม่กรอก Email", "Negative", "R-02, R-08"),
    ("TS-09", "Login", "Login โดยไม่กรอก Password", "Negative", "R-02, R-08"),
    ("TS-10", "Login", "Login ด้วย Email ที่ยังไม่ได้สมัครสมาชิก", "Negative", "R-08"),
]
for s in scenarios:
    ws.append(list(s))
finish(ws, 5, 3 + len(scenarios), [12, 14, 52, 11, 14], title_row=3)

# ---------- Sheet 2: Test Cases ----------
ws = wb.create_sheet("Test Cases")
headers = ["Test Case ID", "Module", "Type", "Test Case", "Test Data", "Steps", "Expected Result", "Actual Result", "Status"]
ws.append(headers)
style_header(ws, len(headers), row=1)
tcs = [
    ("TC-REG-001", "Registration", "Positive", "สมัครสมาชิกด้วยข้อมูลถูกต้องทุกช่อง",
     "Name: Somchai Sukhumvit\nEmail: somchai.test@gmail.com\nPassword: Test1234\nConfirm: Test1234",
     "1) เปิดหน้า Registration\n2) กรอก Name\n3) กรอก Email\n4) กรอก Password\n5) กรอก Confirm Password\n6) กดปุ่ม Register",
     "สมัครสมาชิกสำเร็จ แสดงข้อความยืนยันการสมัคร และนำบัญชีไป Login ได้",
     "สมัครสำเร็จ ขึ้นข้อความ Register Success", "PASS"),
    ("TC-REG-002", "Registration", "Positive", "สมัครสมาชิกบัญชีที่สอง (ผู้ใช้หลายคน)",
     "Name: Mali Jaidee\nEmail: mali.test@gmail.com\nPassword: Test1234\nConfirm: Test1234",
     "1) เปิดหน้า Registration\n2) กรอกข้อมูลบัญชีใหม่ให้ครบ\n3) กดปุ่ม Register",
     "รองรับผู้ใช้หลายบัญชี สมัครบัญชีที่สองสำเร็จ",
     "สมัครสำเร็จ", "PASS"),
    ("TC-REG-003", "Registration", "Negative", "สมัครสมาชิกโดยไม่กรอก Name",
     "Name: (ว่าง)\nEmail: test01@gmail.com\nPassword: Test1234\nConfirm: Test1234",
     "1) เปิดหน้า Registration\n2) เว้นช่อง Name ว่าง\n3) กรอกช่องที่เหลือให้ครบ\n4) กดปุ่ม Register",
     "แสดงข้อความแจ้งเตือน กรุณากรอกชื่อ และไม่สมัครสำเร็จ",
     "แจ้งเตือน Please enter your name ตามที่คาดหวัง", "PASS"),
    ("TC-REG-004", "Registration", "Negative", "สมัครสมาชิกโดยไม่กรอก Email",
     "Name: Somchai\nEmail: (ว่าง)\nPassword: Test1234\nConfirm: Test1234",
     "1) เปิดหน้า Registration\n2) เว้นช่อง Email ว่าง\n3) กรอกช่องที่เหลือให้ครบ\n4) กดปุ่ม Register",
     "แสดงข้อความแจ้งเตือน กรุณากรอก Email และไม่สมัครสำเร็จ",
     "แจ้งเตือน Please enter email ตามที่คาดหวัง", "PASS"),
    ("TC-REG-005", "Registration", "Negative", "สมัครสมาชิกด้วย Email ผิดรูปแบบ (ไม่มี @)",
     "Name: Somchai\nEmail: testgmail.com\nPassword: Test1234\nConfirm: Test1234",
     "1) เปิดหน้า Registration\n2) กรอก Name\n3) กรอก Email เป็น testgmail.com (ไม่มี @)\n4) กรอก Password และ Confirm\n5) กดปุ่ม Register",
     "แสดงข้อความแจ้งเตือนว่ารูปแบบ Email ไม่ถูกต้อง และไม่สมัครสำเร็จ",
     "ระบบไม่แจ้งเตือน และสมัครสมาชิกสำเร็จด้วย Email ที่ผิดรูปแบบ", "FAIL"),
    ("TC-REG-006", "Registration", "Negative", "สมัครสมาชิกที่ Password สั้นกว่า 6 ตัวอักษร",
     "Name: Somchai\nEmail: test02@gmail.com\nPassword: abc1\nConfirm: abc1",
     "1) เปิดหน้า Registration\n2) กรอกข้อมูลครบ โดย Password ใช้ abc1 (4 ตัว)\n3) กดปุ่ม Register",
     "แจ้งเตือน Password ต้องมีอย่างน้อย 6 ตัวอักษร และไม่สมัครสำเร็จ",
     "แจ้งเตือนตามที่คาดหวัง", "PASS"),
    ("TC-REG-007", "Registration", "Negative", "สมัครสมาชิกที่ Password ไม่ตรงกับ Confirm Password",
     "Name: Somchai\nEmail: test03@gmail.com\nPassword: Test1234\nConfirm: Test9999",
     "1) เปิดหน้า Registration\n2) กรอกข้อมูลครบ โดย Confirm Password ต่างจาก Password\n3) กดปุ่ม Register",
     "แจ้งเตือน Password ไม่ตรงกัน และไม่สมัครสำเร็จ",
     "แจ้งเตือน Passwords do not match ตามที่คาดหวัง", "PASS"),
    ("TC-REG-008", "Registration", "Negative", "สมัครสมาชิกด้วย Email ที่มีในระบบแล้ว",
     "Name: Somchai B\nEmail: somchai.test@gmail.com (สมัครแล้วจาก TC-REG-001)\nPassword: Test1234\nConfirm: Test1234",
     "1) เปิดหน้า Registration\n2) กรอกข้อมูลครบโดยใช้ Email ที่สมัครไปแล้ว\n3) กดปุ่ม Register",
     "แจ้งเตือน Email นี้ถูกใช้งานแล้ว และไม่สมัครสำเร็จ",
     "เกิด error และหน้าเว็บว่าง (Blank page) ไม่มีข้อความแจ้งเตือนที่เข้าใจได้", "FAIL"),
    ("TC-LOGIN-001", "Login", "Positive", "Login ด้วย Email และ Password ที่ถูกต้อง",
     "Email: somchai.test@gmail.com\nPassword: Test1234",
     "1) เปิดหน้า Login\n2) กรอก Email\n3) กรอก Password\n4) กดปุ่ม Login",
     "เข้าสู่ระบบสำเร็จ ไปที่หน้าหลัก (Dashboard)",
     "เข้าสู่ระบบสำเร็จ ไปที่หน้า Dashboard", "PASS"),
    ("TC-LOGIN-002", "Login", "Positive", "Login สำเร็จ ออกจากระบบ แล้ว Login ซ้ำได้อีกครั้ง",
     "Email: somchai.test@gmail.com\nPassword: Test1234",
     "1) Login สำเร็จ\n2) กด Logout\n3) Login อีกครั้งด้วยข้อมูลเดิม",
     "Logout ได้ และ Login ซ้ำได้สำเร็จทุกครั้ง",
     "Logout และ Login ซ้ำได้ตามปกติ", "PASS"),
    ("TC-LOGIN-003", "Login", "Positive", "ผู้ใช้ที่เพิ่งสมัครใหม่ Login ได้ทันที",
     "Email: mali.test@gmail.com (เพิ่งสมัครจาก TC-REG-002)\nPassword: Test1234",
     "1) สมัครสมาชิกใหม่ให้สำเร็จ\n2) ไปหน้า Login\n3) Login ด้วยบัญชีที่เพิ่งสมัคร",
     "ใช้บัญชีใหม่ Login ได้ทันทีโดยไม่ต้องรอ",
     "Login สำเร็จทันที", "PASS"),
    ("TC-LOGIN-004", "Login", "Negative", "Login โดยไม่กรอก Email",
     "Email: (ว่าง)\nPassword: Test1234",
     "1) เปิดหน้า Login\n2) เว้นช่อง Email ว่าง\n3) กรอก Password\n4) กดปุ่ม Login",
     "แจ้งเตือน กรุณากรอก Email และไม่เข้าสู่ระบบ",
     "ไม่มีข้อความแจ้งเตือน กดปุ่ม Login แล้วไม่มีอะไรเกิดขึ้น ผู้ใช้งงว่าเกิดอะไรขึ้น", "FAIL"),
    ("TC-LOGIN-005", "Login", "Negative", "Login โดยไม่กรอก Password",
     "Email: somchai.test@gmail.com\nPassword: (ว่าง)",
     "1) เปิดหน้า Login\n2) กรอก Email\n3) เว้นช่อง Password ว่าง\n4) กดปุ่ม Login",
     "แจ้งเตือน กรุณากรอก Password และไม่เข้าสู่ระบบ",
     "แจ้งเตือน Please enter password ตามที่คาดหวัง", "PASS"),
    ("TC-LOGIN-006", "Login", "Negative", "Login ด้วย Password ที่ผิด",
     "Email: somchai.test@gmail.com\nPassword: Wrong999",
     "1) เปิดหน้า Login\n2) กรอก Email ถูกต้อง\n3) กรอก Password ที่ผิด\n4) กดปุ่ม Login",
     "แจ้งเตือน Email หรือ Password ไม่ถูกต้อง และไม่เข้าสู่ระบบ",
     "แจ้งเตือน Invalid email or password ตามที่คาดหวัง", "PASS"),
    ("TC-LOGIN-007", "Login", "Negative", "Login ด้วย Email ที่ยังไม่ได้สมัครสมาชิก",
     "Email: nobody.test@gmail.com (ไม่มีในระบบ)\nPassword: Test1234",
     "1) เปิดหน้า Login\n2) กรอก Email ที่ไม่เคยสมัคร\n3) กรอก Password\n4) กดปุ่ม Login",
     "แจ้งเตือนว่าไม่พบบัญชีผู้ใช้ และไม่เข้าสู่ระบบ",
     "แจ้งเตือน Account not found ตามที่คาดหวัง", "PASS"),
]
for tc in tcs:
    ws.append(list(tc))
finish(ws, 9, 1 + len(tcs), [14, 13, 10, 34, 34, 40, 38, 38, 9])
for r in range(2, 2 + len(tcs)):
    st = ws.cell(row=r, column=9)
    st.alignment = CENTER
    st.font = Font(bold=True, color="006100" if st.value == "PASS" else "9C0006")
    st.fill = PASS_FILL if st.value == "PASS" else FAIL_FILL

# ---------- Sheet 3: Bug Reports ----------
ws = wb.create_sheet("Bug Reports")
headers = ["Bug ID", "Bug Title", "Related Test Case", "Steps to Reproduce", "Expected Result", "Actual Result", "Severity", "Priority", "Status"]
ws.append(headers)
style_header(ws, len(headers), row=1)
bugs = [
    ("BUG-001", "ระบบอนุญาตให้สมัครสมาชิกด้วย Email ที่ไม่ถูกต้อง (ไม่มีเครื่องหมาย @)", "TC-REG-005",
     "1) เปิดหน้า Registration\n2) กรอก Name\n3) กรอก Email เป็น testgmail.com\n4) กรอก Password และ Confirm Password\n5) กดปุ่ม Register",
     "ระบบควรแจ้งเตือนว่ารูปแบบ Email ไม่ถูกต้อง และไม่ให้สมัครสำเร็จ",
     "ระบบไม่แจ้งเตือน และสมัครสมาชิกสำเร็จด้วย Email ที่ผิดรูปแบบ", "Medium", "Medium", "Open"),
    ("BUG-002", "กดปุ่ม Login โดยไม่กรอก Email แล้วไม่มีข้อความแจ้งเตือน และไม่มีอะไรเกิดขึ้น", "TC-LOGIN-004",
     "1) เปิดหน้า Login\n2) เว้นช่อง Email ว่าง\n3) กรอก Password\n4) กดปุ่ม Login",
     "ระบบควรแจ้งเตือน กรุณากรอก Email เพื่อบอกผู้ใช้ว่าข้อมูลยังไม่ครบ",
     "ไม่มีข้อความแจ้งเตือนใด ๆ กดปุ่มแล้วเหมือนไม่มีอะไรเกิดขึ้น ผู้ใช้ไม่ทราบสาเหตุ", "Low", "Low", "Open"),
    ("BUG-003", "สมัครสมาชิกด้วย Email ที่มีในระบบแล้ว ระบบแสดงหน้าเว็บว่าง (Blank Page) แทนข้อความแจ้งเตือน", "TC-REG-008",
     "1) สมัครสมาชิกด้วย Email somchai.test@gmail.com ให้สำเร็จ\n2) สมัครอีกครั้งด้วย Email เดิม\n3) กดปุ่ม Register",
     "ระบบควรแจ้งเตือน Email นี้ถูกใช้งานแล้ว และไม่ให้สมัครซ้ำ",
     "หน้าเว็บกลายเป็นหน้าว่าง (Blank Page) ไม่แสดงข้อความใด ๆ ผู้ใช้ไม่ทราบว่าสมัครสำเร็จหรือไม่", "High", "High", "Open"),
]
for b in bugs:
    ws.append(list(b))
finish(ws, 9, 1 + len(bugs), [10, 44, 14, 44, 38, 42, 10, 10, 9])
for r in range(2, 2 + len(bugs)):
    sev = ws.cell(row=r, column=7)
    sev.alignment = CENTER
    sev.font = Font(bold=True)
    if sev.value in SEV_FILL:
        sev.fill = SEV_FILL[sev.value]

# ---------- Sheet 4: Test Summary ----------
ws = wb.create_sheet("Test Summary")
ws["A1"] = "Login & Registration Testing — Test Summary Report"
ws["A1"].font = TITLE_FONT
ws.append([])
rows = [
    ("รายการ", "จำนวน"),
    ("Test Scenarios", 10),
    ("Test Cases ทั้งหมด", 15),
    ("ผ่าน (Passed)", 12),
    ("ไม่ผ่าน (Failed)", 3),
    ("Pass Rate", "80%"),
    ("Bugs ที่พบ", "3 (High 1 / Medium 1 / Low 1)"),
]
for row in rows:
    ws.append(list(row))
style_header(ws, 2, row=3)
ws.append([])
ws.append(["สรุปผลการทดสอบ"])
ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=12, color="1F4E79")
ws.append([])
ws.append(["ระบบสมัครสมาชิกและเข้าสู่ระบบ (Login & Registration) สามารถทำงานตาม Requirement ส่วนใหญ่ได้ ฟังก์ชันหลักใช้งานได้ปกติ (Positive Tests ผ่านทั้งหมด 5/5) และจัดการข้อมูลผิดพลาดได้ส่วนใหญ่ (Negative Tests ผ่าน 7/10) แต่ยังพบปัญหา 3 จุดที่ควรแก้ไขก่อนนำระบบไปใช้งานจริง ได้แก่ BUG-001 (สมัครได้ด้วย Email ผิดรูปแบบ) BUG-002 (ไม่แจ้งเตือนเมื่อไม่กรอก Email ตอน Login) และ BUG-003 (หน้าเว็บว่างเมื่อสมัครด้วย Email ซ้ำ ระดับ High)"])
ws.cell(row=ws.max_row, column=1).alignment = WRAP
ws.append([])
ws.append(["คำแนะนำ: แก้ไข BUG-003 (High) เป็นอันดับแรก จากนั้น BUG-001 และ BUG-002 ตามลำดับ หลังแก้ไขแล้วทดสอบซ้ำ (Regression Testing) เคสที่เคย FAIL ทั้ง 3 เคส"])
ws.cell(row=ws.max_row, column=1).alignment = WRAP
ws.column_dimensions["A"].width = 40
ws.column_dimensions["B"].width = 30
ws.merge_cells(start_row=11, start_column=1, end_row=11, end_column=2)
ws.merge_cells(start_row=13, start_column=1, end_row=13, end_column=2)
ws.row_dimensions[11].height = 90
ws.row_dimensions[13].height = 45

wb.save(OUT)
print("saved:", OUT)
